import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path
import os
import shutil

INPUT_DOCX = "/Users/hirunbhathiya/Downloads/Intelligent_Resource_Allocation_Productivity_Optimizer_Thesis_Draft.docx"
OUTPUT_DOCX = "/Users/hirunbhathiya/Downloads/Intelligent_Resource_Allocation_Productivity_Optimizer_Thesis_Final.docx"
FIGURES_DIR = Path("/Users/hirunbhathiya/Desktop/RP module/R26-SE-005/ai-service/models/graphs/thesis_figures")

def set_run_font(run, font_name="Times New Roman", size_pt=11, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

def set_paragraph_format(p, font_name="Times New Roman", size_pt=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    for r in p.runs:
        set_run_font(r, font_name, size_pt, bold, italic)

def set_cell_properties(cell, fill_hex=None, top_pad=100, bottom_pad=100, left_pad=150, right_pad=150):
    tcPr = cell._tc.get_or_add_tcPr()
    if fill_hex:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top_pad}" w:type="dxa"/><w:bottom w:w="{bottom_pad}" w:type="dxa"/><w:left w:w="{left_pad}" w:type="dxa"/><w:right w:w="{right_pad}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_cell(cell, text, bold=False, italic=False, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", font_size, bold, italic)
    if fill_hex:
        set_cell_properties(cell, fill_hex=fill_hex)

def create_table_after_paragraph(doc, p, headers, data, col_widths=None, table_title=None):
    """Creates a table, formats it, and inserts it right after paragraph p."""
    # If table title provided, insert title paragraph
    current_elem = p._p
    if table_title:
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = title_p.add_run(table_title)
        set_run_font(r, "Times New Roman", 10.5, bold=True)
        current_elem.addnext(title_p._p)
        current_elem = title_p._p

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        format_cell(hdr_cells[i], title, bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex="E2E8F0")

    # Format Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        fill = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            format_cell(row_cells[c_idx], str(val), bold=False, font_size=9.0, fill_hex=fill)

    # Set column widths if specified
    if col_widths:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)

    current_elem.addnext(table._tbl)
    return table

def insert_image_and_caption_after_paragraph(doc, p, img_path, caption_text, width_inches=5.8):
    """Inserts image paragraph and caption paragraph right after paragraph p."""
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(str(img_path), width=Inches(width_inches))

    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_r = cap_p.add_run(caption_text)
    set_run_font(cap_r, "Times New Roman", 9.5, bold=True, italic=True)

    # Position in XML
    p._p.addnext(cap_p._p)
    p._p.addnext(img_p._p)

def replace_placeholder_with_image_and_caption(doc, p, img_path, caption_text, width_inches=5.8):
    """Clears placeholder text in p, embeds image into p, and adds caption immediately after."""
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(width_inches))

    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_r = cap_p.add_run(caption_text)
    set_run_font(cap_r, "Times New Roman", 9.5, bold=True, italic=True)

    p._p.addnext(cap_p._p)

print("Setup helper functions complete")
